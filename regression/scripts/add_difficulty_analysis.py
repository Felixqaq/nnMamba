#!/usr/bin/env python3
"""Add the difficulty-stratified breakdown to the doctor's workbook.

A patient whose FEV1/FVC sits a point either side of 70 is labelled Normal or
Abnormal by a hair, while the lungs look the same. Reporting one pooled AUC over
a cohort that is 45% such patients hides where the model does and does not work,
so this writes the stratification alongside the pooled number and tags every
validation patient with its distance from the threshold.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    confusion_matrix,
    roc_auc_score,
)

BANDS = [
    ("全部 200 人", None, None),
    ("邊界 |比值-70| < 5", None, 5.0),
    ("邊界 |比值-70| < 7", None, 7.0),
    ("邊界 |比值-70| < 10", None, 10.0),
    ("明確 |比值-70| >= 5", 5.0, None),
    ("明確 |比值-70| >= 7", 7.0, None),
    ("明確 |比值-70| >= 10", 10.0, None),
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--excel", type=Path, required=True)
    p.add_argument("--mamba-json", type=Path, required=True)
    p.add_argument("--tapct-json", type=Path, required=True)
    p.add_argument("--pft-csv", type=Path, required=True)
    p.add_argument("--bootstrap", type=int, default=5000)
    p.add_argument("--seed", type=int, default=20260901)
    return p.parse_args()


def load_predictions(path: Path):
    """Return per-patient probability, truth and the decision already made.

    The Mamba file's headline metrics come from a majority vote of five members,
    each thresholded at the frozen training-OOF value -- not from cutting the
    mean probability. Re-thresholding the mean here would silently produce
    different accuracy/sensitivity than the run reported, so take pred_label as
    written and keep the probability only for AUC, which is threshold-free.
    """
    d = json.loads(path.read_text(encoding="utf-8-sig"))
    if "patients" not in d:
        raise SystemExit(f"{path}: no 'patients' block; keys={sorted(d)}")
    if not d["patients"]:
        raise SystemExit(f"{path}: 'patients' is empty")
    prob, truth, decided = {}, {}, {}
    for pid, row in d["patients"].items():
        value = row.get("mean_prob_abnormal", row.get("prob_abnormal"))
        label = row.get("true_label")
        if value is None or label is None:
            continue
        prob[pid] = float(value)
        truth[pid] = 1 if label == "Abnormal" else 0
        pred = row.get("pred_label")
        if pred is not None:
            decided[pid] = 1 if pred == "Abnormal" else 0
    if len(decided) != len(prob):
        raise SystemExit(f"{path}: not every patient carries a pred_label")
    return prob, truth, decided


def boot_auc_ci(y: np.ndarray, p: np.ndarray, n: int, seed: int) -> tuple[float, float]:
    rng = np.random.default_rng(seed)
    vals = []
    for _ in range(n):
        i = rng.integers(0, len(y), len(y))
        if len(set(y[i].tolist())) < 2:
            continue
        vals.append(roc_auc_score(y[i], p[i]))
    if not vals:
        return float("nan"), float("nan")
    return float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5))


def band_rows(tag, prob, truth, decided, ratio, boot, seed):
    rows = []
    for name, lo, hi in BANDS:
        ids = []
        for pid in prob:
            if pid not in ratio:
                continue
            d = abs(ratio[pid] - 70.0)
            if hi is not None and not d < hi:
                continue
            if lo is not None and not d >= lo:
                continue
            ids.append(pid)
        y = np.array([truth[p] for p in ids], dtype=int)
        s = np.array([prob[p] for p in ids], dtype=float)
        if len(ids) == 0 or len(set(y.tolist())) < 2:
            rows.append([tag, name, len(ids), int(y.sum()) if len(ids) else 0,
                         "—", "—", "—", "—", "—", "—", "類別單一,無法計算"])
            continue
        pred = np.array([decided[p] for p in ids], dtype=int)
        auc = roc_auc_score(y, s)
        ci_lo, ci_hi = boot_auc_ci(y, s, boot, seed)
        tn, fp, fn, tp = confusion_matrix(y, pred, labels=[0, 1]).ravel()
        note = "信賴區間涵蓋 0.5,與擲骰子無異" if ci_lo <= 0.5 else ""
        rows.append([
            tag, name, len(ids), int(y.sum()),
            round(auc, 4), f"[{ci_lo:.4f}, {ci_hi:.4f}]",
            round(float(accuracy_score(y, pred)), 4),
            round(float(balanced_accuracy_score(y, pred)), 4),
            round(float(tp / (tp + fn)) if (tp + fn) else float("nan"), 4),
            round(float(tn / (tn + fp)) if (tn + fp) else float("nan"), 4),
            note,
        ])
    return rows


def style(sheet, table_name: str) -> None:
    sheet.freeze_panes = "A2"
    # No sheet.auto_filter here: the Table below carries its own <autoFilter>,
    # and setting both makes Excel declare the workbook corrupt on open.
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if sheet.max_row >= 2:
        t = Table(displayName=table_name, ref=sheet.dimensions)
        t.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        sheet.add_table(t)
    for column in sheet.columns:
        width = min(46, max(11, max(len(str(c.value or "")) for c in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def main() -> None:
    args = parse_args()

    ratio: dict[str, float] = {}
    with io.open(args.pft_csv, encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        # The header of fev1_fvc.csv has shipped padded before (" PatientID"),
        # so names are stripped -- but a *missing* column must be fatal. Reading
        # it with a "" default would leave every ratio absent and quietly emit a
        # difficulty table where every band has n=0, which looks like a result.
        fields = {(f or "").strip() for f in (reader.fieldnames or [])}
        for required in ("PatientID", "FEV1FVC_pct"):
            if required not in fields:
                raise SystemExit(
                    f"{args.pft_csv}: no {required!r} column; found {sorted(fields)[:8]}"
                )
        for row in reader:
            row = {(k or "").strip(): (v or "").strip() for k, v in row.items()}
            value = row["FEV1FVC_pct"]
            if value:
                try:
                    ratio[row["PatientID"]] = float(value)
                except ValueError:
                    pass
    if not ratio:
        raise SystemExit(f"{args.pft_csv}: parsed 0 usable FEV1/FVC values")

    mp, mt, md = load_predictions(args.mamba_json)
    tp_, tt, td = load_predictions(args.tapct_json)
    print(f"Mamba n={len(mp)}  已判定 {sum(md.values())} 位為 Abnormal")
    print(f"TapCT n={len(tp_)}  已判定 {sum(td.values())} 位為 Abnormal")

    uncovered = sorted(set(mp) - set(ratio))
    if uncovered:
        raise SystemExit(
            f"{len(uncovered)} predicted patients have no FEV1/FVC in the CSV, so "
            f"their difficulty band is unknown: {uncovered[:5]}"
        )

    rows = (band_rows("Mamba5", mp, mt, md, ratio, args.bootstrap, args.seed)
            + band_rows("TAPCT", tp_, tt, td, ratio, args.bootstrap, args.seed))

    wb = load_workbook(args.excel)

    # ---- difficulty columns on the validation sheet ----
    ws = wb["Validation_200"]
    headers = [c.value for c in ws[1]]
    if "Distance_from_70" not in headers:
        # Insert before Batch rather than appending, so Batch stays the last
        # column of the sheet as the workbook is laid out elsewhere.
        col = (headers.index("Batch") + 1) if "Batch" in headers else ws.max_column + 1
        ws.insert_cols(col, 2)
        ws.cell(row=1, column=col, value="Distance_from_70")
        ws.cell(row=1, column=col + 1, value="Difficulty_Band")
        headers = [c.value for c in ws[1]]
        pid_col = headers.index("Patient_ID") + 1
        for r in range(2, ws.max_row + 1):
            pid = str(ws.cell(row=r, column=pid_col).value)
            if pid in ratio:
                d = abs(ratio[pid] - 70.0)
                ws.cell(row=r, column=col, value=round(d, 1))
                ws.cell(row=r, column=col + 1,
                        value="邊界 (<7)" if d < 7 else "明確 (>=7)")
        # The existing table must grow to cover the new columns. In this
        # openpyxl version ws.tables maps name -> ref string, not to the Table
        # object, so read the ref straight from the mapping.
        for name, ref in list(ws.tables.items()):
            if not isinstance(ref, str):
                ref = ref.ref
            del ws.tables[name]
            start, end = ref.split(":")
            end_col = ws.cell(row=1, column=ws.max_column).column_letter
            end_row = "".join(ch for ch in end if ch.isdigit())
            new = Table(displayName=name, ref=f"{start}:{end_col}{end_row}")
            new.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
            )
            ws.add_table(new)
        for c in (col, col + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[cell.column_letter].width = 18
        print("added Distance_from_70 / Difficulty_Band to Validation_200")

    # ---- the analysis sheet ----
    if "Difficulty_Analysis" in wb.sheetnames:
        del wb["Difficulty_Analysis"]
    sheet = wb.create_sheet("Difficulty_Analysis")
    sheet.append(["模型", "族群", "人數", "異常數", "AUC", "AUC 95% CI",
                  "Accuracy", "Balanced_Acc", "敏感度", "特異度", "備註"])
    for row in rows:
        sheet.append(row)
    style(sheet, "DifficultyAnalysis")

    wb.save(args.excel)
    print(f"wrote {args.excel}")
    print("\n{:8s}{:24s}{:>6}{:>7}{:>9}  {}".format(
        "模型", "族群", "n", "異常", "AUC", "95% CI"))
    for r in rows:
        print("{:8s}{:24s}{:>6}{:>7}{:>9}  {}".format(
            r[0], r[1], r[2], r[3], str(r[4]), str(r[5])))


if __name__ == "__main__":
    main()

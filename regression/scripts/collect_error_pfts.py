#!/usr/bin/env python3
"""Collect the spirometry reports of every misclassified holdout patient.

A clinician looking at these is usually asking whether the *label* is sound --
an under-blown manoeuvre biases FEV1/FVC upward and can put an obstructed
patient on the Normal side of 70. So the files are grouped by which model erred,
and each filename carries the ratio and the confusion class, with a plain-language
key written beside the images because FP/FN is not everyone's shorthand.
"""

from __future__ import annotations

import argparse
import csv
import glob
import io
import json
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

GROUPS = {
    "both_wrong": "兩個模型都判錯",
    "mamba_only": "只有 Mamba 判錯",
    "tapct_only": "只有 TAPCT 判錯",
}

README_LINES = [
    "誤判病人的肺功能報告",
    "====================",
    "",
    "檔名格式",
    "  <病人編號>_fev1fvc_<實測FEV1/FVC百分比>_<誤判型態>.jpg",
    "  例:A123456_fev1fvc_71_FP.jpg",
    "",
    "誤判型態的兩個字母",
    "  第二個字母 = 模型判成什麼",
    "     P = Positive = 陽性 = Abnormal(阻塞)",
    "     N = Negative = 陰性 = Normal(正常)",
    "  第一個字母 = 判得對不對",
    "     T = True  = 判對了",
    "     F = False = 判錯了",
    "",
    "  本資料夾只收判錯的病人,所以只會出現這兩種:",
    "     FP  偽陽性:實際是 Normal,被判成 Abnormal   -> 誤報",
    "     FN  偽陰性:實際是 Abnormal,被判成 Normal   -> 漏診",
    "  判對的 TP / TN 不在這個資料夾裡。",
    "",
    "標籤規則",
    "  FEV1/FVC < 70% = Abnormal(嚴格小於,剛好 70 算 Normal)",
    "",
    "資料夾",
    "  both_wrong   Mamba 與 TAPCT 都判錯",
    "  mamba_only   只有 Mamba 判錯,TAPCT 判對",
    "  tapct_only   只有 TAPCT 判錯,Mamba 判對",
    "",
    "請醫師特別留意",
    "  比值落在 70 附近的病人,如果吹氣不確實,FVC 被截短的程度會大於 FEV1,",
    "  使 FEV1/FVC 被高估 —— 一位真正阻塞的病人就可能被記錄成正常。",
    "  這類個案若判定為檢查無效,可以從世代中排除,",
    "  如同先前因吹氣無效而被排除的那兩位。",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--mamba-json", type=Path, required=True)
    p.add_argument("--tapct-json", type=Path, required=True)
    p.add_argument("--pft-csv", type=Path, required=True)
    p.add_argument("--jpg-root", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--margin", type=float, default=7.0)
    return p.parse_args()


def load_clinical(path: Path) -> dict[str, dict[str, str]]:
    rows: dict[str, dict[str, str]] = {}
    with io.open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        fields = {(f or "").strip() for f in (reader.fieldnames or [])}
        for required in ("PatientID", "FEV1FVC_pct"):
            if required not in fields:
                raise SystemExit(f"{path}: no {required!r} column")
        for row in reader:
            row = {(k or "").strip(): (v or "").strip() for k, v in row.items()}
            rows[row["PatientID"]] = row
    if not rows:
        raise SystemExit(f"{path}: no rows")
    return rows


def main() -> None:
    args = parse_args()
    clinical = load_clinical(args.pft_csv)

    mamba = json.loads(args.mamba_json.read_text(encoding="utf-8-sig"))["patients"]
    tapct = json.loads(args.tapct_json.read_text(encoding="utf-8-sig"))["patients"]

    records = []
    for pid, row in mamba.items():
        truth = row["true_label"]
        m_wrong = row["pred_label"] != truth
        t_wrong = tapct[pid]["pred_label"] != truth
        if not (m_wrong or t_wrong):
            continue
        group = ("both_wrong" if m_wrong and t_wrong
                 else "mamba_only" if m_wrong else "tapct_only")
        info = clinical.get(pid)
        if info is None:
            raise SystemExit(f"{pid}: no clinical row in {args.pft_csv}")
        ratio = float(info["FEV1FVC_pct"])
        distance = abs(ratio - 70.0)
        # Every patient here is an error, so the erring model's class is either
        # FP (called Abnormal, is Normal) or FN (called Normal, is Abnormal).
        records.append({
            "pid": pid,
            "group": group,
            "truth": truth,
            "confusion": "FP" if truth == "Normal" else "FN",
            "mamba": row["pred_label"],
            "votes": row.get("vote_text", ""),
            "tapct": tapct[pid]["pred_label"],
            "tapct_prob": tapct[pid].get("prob_abnormal"),
            "ratio": ratio,
            "distance": round(distance, 1),
            "band": "邊界" if distance < args.margin else "明確",
            "sex": info.get("Sex", ""),
            "age": int(float(info["Age"])) if info.get("Age") else None,
            "fev1_ref": info.get("FEV1_pctpred", ""),
            "batch": info.get("Date", ""),
        })

    if args.out.exists():
        shutil.rmtree(args.out)
    for name in GROUPS:
        (args.out / name).mkdir(parents=True)

    copied, missing = 0, []
    for r in records:
        hits = glob.glob(str(args.jpg_root / "*" / f"{r['pid']}.jpg"))
        if not hits:
            missing.append(r["pid"])
            continue
        name = f"{r['pid']}_fev1fvc_{r['ratio']:.0f}_{r['confusion']}.jpg"
        shutil.copy2(hits[0], args.out / r["group"] / name)
        copied += 1
    if missing:
        raise SystemExit(f"no PFT jpg for {len(missing)} patients: {missing[:6]}")

    (args.out / "檔名說明.txt").write_text(
        "\n".join(README_LINES) + "\n", encoding="utf-8"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "誤判病人"
    ws.append(["病人編號", "分組", "誤判型態", "誤判型態說明", "真實標籤",
               "Mamba判定", "Mamba票數", "TAPCT判定", "TAPCT機率",
               "FEV1_FVC_pct", "距離70", "難度", "性別", "年齡",
               "FEV1_REF_pct", "檢查批次"])
    order = {"both_wrong": 0, "mamba_only": 1, "tapct_only": 2}
    for r in sorted(records, key=lambda x: (order[x["group"]], x["distance"])):
        explain = ("FP 偽陽性:實際 Normal,被判成 Abnormal(誤報)"
                   if r["confusion"] == "FP"
                   else "FN 偽陰性:實際 Abnormal,被判成 Normal(漏診)")
        ws.append([r["pid"], GROUPS[r["group"]], r["confusion"], explain, r["truth"],
                   r["mamba"], r["votes"], r["tapct"], r["tapct_prob"],
                   r["ratio"], r["distance"], r["band"], r["sex"], r["age"],
                   r["fev1_ref"], r["batch"]])

    ws.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    # Table only, never sheet.auto_filter as well: two filter declarations over
    # one range make Excel declare the workbook corrupt on open.
    table = Table(displayName="ErrorCases", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    for column in ws.columns:
        width = min(46, max(11, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    wb.save(args.out / "誤判病人索引.xlsx")

    print(f"複製了 {copied} 張 PFT 報告到 {args.out}")
    for name, label in GROUPS.items():
        n = len(list((args.out / name).glob("*.jpg")))
        fp = sum(1 for r in records if r["group"] == name and r["confusion"] == "FP")
        fn = sum(1 for r in records if r["group"] == name and r["confusion"] == "FN")
        border = sum(1 for r in records if r["group"] == name and r["band"] == "邊界")
        print(f"  {name:12s} {label:16s} {n:>3} 張  FP {fp:>2} / FN {fn:>2}  (邊界 {border})")
    total_border = sum(1 for r in records if r["band"] == "邊界")
    print(f"\n  {len(records)} 位誤判者中,{total_border} 位是邊界病例 "
          f"({total_border / len(records):.0%})")


if __name__ == "__main__":
    main()
